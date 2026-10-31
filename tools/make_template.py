"""The show-script template, as .docx (opens in Google Docs and Word) and .md."""
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
import os as _os
ROOT = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))   # the kit (this repo)
BUILD = _os.path.join(ROOT, "tools", ".build"); _os.makedirs(BUILD, exist_ok=True)
KIT = _os.path.join(ROOT, "templates")
IMG = sys.argv[1] if len(sys.argv) > 1 else _os.path.join(ROOT, "icons", "hotkeys", "12-winner.png")   # the example photo

# (kind, text) — kinds: title, note, h1, intro, h2, h3, p (with **bold**, *italic*, __underline__), bullet, number, hr, img
S = [
 ("title", "Argue with Yo Momma! · Episode 12 script"),
 ("note", "Note: How this works: each Heading 1 is a segment, each Heading 2 is one on-screen card (its heading is the card's title). Paragraphs that start with \"Note:\" (like this one) never go on screen. Import the doc in the Content Panel → Import a script."),
 ("h1", "Opening Tip-Off (3 min)"),
 ("intro", "Welcome in: who's here and what's on the table"),
 ("h2", "Tonight's lineup"),
 ("bullet", "**Hot Takes**: say it with your chest"),
 ("bullet", "**Buy or Sell**: the league's hottest stocks"),
 ("bullet", "**Tin Foil Hat**: crazy or credible?"),
 ("h1", "Hot Takes (5 min)"),
 ("intro", "Say it with your chest"),
 ("h2", "Take 1"),
 ("p", "**The Knicks win the East this year.** Not *maybe*. __This year.__"),
 ("note", "Note: Mom goes first on this one."),
 ("h2", "Take 2"),
 ("p", "The best player in the league isn't on a contender."),
 ("img", ""),
 ("h1", "Buy or Sell (5 min)"),
 ("intro", "Quick verdicts on the league's hottest stocks"),
 ("h2", "Buy or sell?"),
 ("number", "Rookie of the Year race is already over"),
 ("number", "The Celtics' bench is a problem"),
 ("number", "Load management is ruining the regular season"),
 ("h1", "Tin Foil Hat (4 min)"),
 ("intro", "Conspiracy theories: crazy or credible?"),
 ("h2", "Theory of the week"),
 ("p", "The draft lottery is *rigged*."),
 ("h3", "Exhibit A"),
 ("p", "Every year the big market team jumps up. Coincidence?"),
 ("hr", ""),
 ("p", "Crazy or credible? Vote in the chat!"),
]

def runs(par, text):
    import re
    for tok in re.split(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*)", text):
        if not tok: continue
        if tok.startswith("**"): par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("__"): par.add_run(tok[2:-2]).underline = True
        elif tok.startswith("*"): par.add_run(tok[1:-1]).italic = True
        else: par.add_run(tok)

d = Document()
for kind, text in S:
    if kind == "title": d.add_heading(text, 0)
    elif kind == "h1": d.add_heading(text, 1)
    elif kind == "h2": d.add_heading(text, 2)
    elif kind == "h3": d.add_heading(text, 3)
    elif kind == "bullet": runs(d.add_paragraph(style="List Bullet"), text)
    elif kind == "number": runs(d.add_paragraph(style="List Number"), text)
    elif kind == "intro": p = d.add_paragraph(); r = p.add_run(text); r.italic = True
    elif kind == "note": p = d.add_paragraph(); r = p.add_run(text); r.font.color.rgb = RGBColor(0x88, 0x88, 0x88); r.font.size = Pt(9)
    elif kind == "hr": d.add_paragraph("―" * 20)  # visual divider in Word; Markdown gets ---
    elif kind == "img":
        if IMG: d.add_picture(IMG, width=Inches(2))
    else: runs(d.add_paragraph(), text)
d.save(KIT + "/AWYM-Show-Script-Template.docx")

md = []
for kind, text in S:
    if kind == "title": md.append(text + "\n")
    elif kind in ("h1", "h2", "h3"): md.append("#" * int(kind[1]) + " " + text + "\n")
    elif kind == "bullet": md.append("- " + text)
    elif kind == "number": md.append("1. " + text)
    elif kind == "intro": md.append("*" + text + "*\n")
    elif kind == "hr": md.append("\n---\n")
    elif kind == "img": continue
    else:
        import re
        md.append(re.sub(r"__([^_]+)__", r"<u>\1</u>", text) + "\n")
open(KIT + "/AWYM-Show-Script-Template.md", "w", encoding="utf-8").write("\n".join(md).replace("\n- ", "\n- ").strip() + "\n")
print("ok")
